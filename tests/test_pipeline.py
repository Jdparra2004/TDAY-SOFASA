"""
Tests del pipeline completo (sin números fijos del universo de prueba).
Los esperados se calculan con código independiente: se lee el Tracking directo con openpyxl.
Ejecutar:  pytest -v tests/ | tee tests/salida_tests.txt
"""
from __future__ import annotations

import inspect
import re
from collections import Counter
from datetime import datetime
from pathlib import Path
import sys

ROOT = Path(__file__).resolve().parent.parent
sys.path.append(str(ROOT))

import numpy as np
import openpyxl
import pandas as pd
import pytest

import main as main_mod
from src.costs import weighted_objective
from src.glossary import get_glossary
from src.model import Params
from src.pipeline import build_params, prepare_universe, print_validaciones, run_pipeline

DATA = ROOT / "data"
FECHA, HORIZONTE, CAP_S, CAP_P, CAP_T, DIAS_T = "2026-10-20", 20, 14, 15, 2, [0, 1, 2, 3, 4]
BASE = pd.Timestamp(FECHA)


def _fecha(v):
    if isinstance(v, datetime):
        return pd.Timestamp(v).normalize() if v.year >= 2020 else None
    return None


def leer_tracking_independiente():
    """Lee el Tracking con openpyxl, sin usar src/: id, ETA (real si existe, si no inicial) y EM real."""
    path = DATA / "1st Data" / "Tracking contenedores.xlsx"
    ws = openpyxl.load_workbook(path, data_only=True).active
    hdr_row = next(i for i, r in enumerate(ws.iter_rows(values_only=True), 1) if r and any(str(c).strip() == "Contenedor" for c in r if c))
    hdr = [str(c).strip().replace("\n", " ") if c else "" for c in next(ws.iter_rows(min_row=hdr_row, max_row=hdr_row, values_only=True))]
    ix = {h: i for i, h in enumerate(hdr)}
    rows, seen = {}, set()
    for r in ws.iter_rows(min_row=hdr_row + 1, values_only=True):
        cid = r[ix["Contenedor"]]
        if not cid or str(cid).strip() in ("", "None"):
            continue
        cid = str(cid).strip()
        if cid in seen:
            continue
        seen.add(cid)
        eta = _fecha(r[ix["ETA Real"]]) or _fecha(r[ix["ETA Inicial"]])
        rows[cid] = {"eta": eta, "em": _fecha(r[ix["EM"]])}
    return rows


@pytest.fixture(scope="module")
def corrida(tmp_path_factory):
    out = tmp_path_factory.mktemp("salida")
    res = main_mod.main(FECHA, HORIZONTE, CAP_S, CAP_P, CAP_T, DIAS_T, data_dir=str(DATA / "1st Data"), out_dir=str(out),
                        time_limit_s=20, workers=8, seed=1)
    assert res is not None
    return res, out


@pytest.fixture(scope="module")
def tracking():
    return leer_tracking_independiente()


def test_universo_embudo_y_ids(corrida, tracking):
    res, out = corrida
    e = res.embudo
    df = res.df_universo
    assert df["id"].is_unique and set(df["id"]).issubset(set(tracking))
    assert e["excluidos"] + e["fuera_ventana"] + e["programables"] == e["tracking_total"] == len(df)
    xl = pd.read_excel(out / "plan_prescriptivo_sofasa.xlsx", sheet_name=None)
    assert len(xl["Excluidos (ya en planta)"]) == e["excluidos"]
    assert len(xl["Fuera de ventana"]) == e["fuera_ventana"]
    assert len(xl["Programacion semanal"]) == e["programables"]


def test_solo_excluye_em_real_y_ventana(corrida, tracking):
    res, _ = corrida
    df = res.df_universo.set_index("id")
    esperados_exc = {i for i, r in tracking.items() if r["em"] is not None and r["em"] < BASE and i in df.index}
    assert set(df.index[df["estado_universo"] == "EXCLUIDO_EN_PLANTA"]) == esperados_exc
    prog = df[df["estado_universo"] == "PROGRAMABLE"]
    assert (prog["eta_dia"] <= HORIZONTE - 1).all()
    fuera = df[df["estado_universo"] == "FUERA_VENTANA"]
    assert (fuera["eta_dia"] > HORIZONTE - 1).all()


def test_eta_coincide_con_tracking(corrida, tracking):
    res, _ = corrida
    malos = [c.id for c in res.containers if tracking[c.id]["eta"] is not None
             and (BASE + pd.Timedelta(days=c.eta)) != tracking[c.id]["eta"]]
    assert not malos, f"ETA distinta al Tracking en {len(malos)} contenedores: {malos[:5]}"


@pytest.mark.parametrize("archivo", ["contenedores_baseline.csv", "contenedores_optimizado.csv"])
def test_capacidades_y_lead_time(corrida, archivo):
    res, out = corrida
    p = res.params
    df = pd.read_csv(out / archivo)
    att = df[df["atendido"]]
    wd = lambda t: (BASE.weekday() + t) % 7
    por_dia = Counter(att["em"].astype(int))
    assert max(por_dia.values(), default=0) <= CAP_S
    sup = Counter((BASE.weekday() + int(t)) // 7 for t in att["suppla_in"].dropna())
    assert max(sup.values(), default=0) <= CAP_P
    tr = Counter(att.loc[att["route"] == "suppla", "em"].astype(int))
    assert max(tr.values(), default=0) <= CAP_T
    assert all(wd(t) in DIAS_T for t in tr)
    assert ((att["em"] - att["eta"]) >= p.l_min).all()


def test_estados_y_nada_desaparece(corrida):
    res, out = corrida
    assert res.estados[["estado_base", "estado_optimo"]].notna().all().all()
    assert set(res.estados["estado_optimo"]) <= {"PROGRAMADO_EN_HORIZONTE", "EM_PROYECTADA", "DIFERIDO_SIN_CUPO"}
    b = pd.read_csv(out / "contenedores_baseline.csv")
    o = pd.read_csv(out / "contenedores_optimizado.csv")
    assert len(b) == len(o) == len(res.containers) == res.embudo["programables"]
    assert set(b["id"]) == set(o["id"])


def test_alertas_en_programacion(corrida):
    res, out = corrida
    prog = pd.read_excel(out / "plan_prescriptivo_sofasa.xlsx", sheet_name="Programacion semanal")
    p = res.params
    lt = pd.to_numeric(prog["Lead time"], errors="coerce")
    assert prog.loc[lt > p.l_max, "Alertas"].fillna("").str.contains("Crítica").all()
    assert prog.loc[(lt > p.l_obj) & (lt <= p.l_max), "Alertas"].fillna("").str.contains("Preventiva").all()


def test_rupturas_solo_con_fecha_real(corrida):
    res, out = corrida
    for archivo in ("contenedores_baseline.csv", "contenedores_optimizado.csv"):
        df = pd.read_csv(out / archivo)
        sin = df[df["need_by"].isna()]
        assert (sin["rupture_days"] == 0).all()
        assert not sin["tiene_ruptura_real"].any()
    assert all(c.need_by is None for c in res.containers if c.ruptura_real is None)
    assert res.base.kpis["n_rupturas"] <= res.base.kpis["n_con_ruptura_real"]
    assert res.optimo.kpis["n_rupturas"] <= res.optimo.kpis["n_con_ruptura_real"]


@pytest.mark.parametrize("archivo,columna", [("contenedores_baseline.csv", "Línea base"), ("contenedores_optimizado.csv", "Plan optimizado")])
def test_costos_recalculados_desde_csv(corrida, archivo, columna):
    res, out = corrida
    p = Params()
    df = pd.read_csv(out / archivo)
    dias_puerto = df["salida_puerto"] - df["eta"]
    assert (dias_puerto == df["dias_puerto"]).all()
    dem_dias = (dias_puerto - p.l_max).clip(lower=0)
    assert (dem_dias == df["dias_demurrage"]).all()
    usd_dem = dem_dias * p.USD_DEMURRAGE_DIA
    suppla_dias = pd.to_numeric(df["em"] - df["suppla_in"], errors="coerce").fillna(0).clip(lower=0)
    expected_suppla = suppla_dias * p.USD_SUPPLA_DIA
    assert np.allclose(df["usd_suppla"], expected_suppla, atol=1.0)
    total = usd_dem.sum() + df["usd_suppla"].sum()
    imp = pd.read_excel(out / "plan_prescriptivo_sofasa.xlsx", sheet_name="Impacto económico")
    fila = imp[(imp["Bloque"] == "Total") & (imp["Concepto"] == "Total (USD)")].iloc[0]
    assert abs(float(fila[columna]) - total) < 1.0
    ah = float(fila["Ahorro (base − óptimo)"])
    assert abs(ah - (float(fila["Línea base"]) - float(fila["Plan optimizado"]))) < 1.0
    prog = pd.read_excel(out / "plan_prescriptivo_sofasa.xlsx", sheet_name="Programacion semanal")
    assert abs(prog["USD demurrage"].sum() + prog["USD almacenamiento Suppla"].sum() - float(fila["Plan optimizado"])) < 1.0


def test_sin_inventario_ni_consumo_y_fechas_reales(corrida):
    res, out = corrida
    for archivo in ("contenedores_baseline.csv", "contenedores_optimizado.csv", "diario_baseline.csv", "diario_optimizado.csv"):
        cols = " ".join(pd.read_csv(out / archivo, nrows=1).columns).lower()
        assert "consumo" not in cols and "inv_sofasa" not in cols and "inventario" not in cols
    prog = pd.read_excel(out / "plan_prescriptivo_sofasa.xlsx", sheet_name="Programacion semanal")
    fechas = prog[["ETD", "ETA", "Ruptura", "EM objetivo", "EM previa", "EM plan"]].astype(str)
    patron = re.compile(r"^(\d{4}-\d{2}-\d{2}|N/D)$")
    assert fechas.apply(lambda s: s.map(lambda v: bool(patron.match(v)))).all().all()
    assert not prog.isin(["None", "NaT"]).any().any()


def test_n_containers_en_kpis_y_csv(corrida):
    res, out = corrida
    k = pd.read_csv(out / "comparacion_kpis.csv", index_col=0)
    assert k.loc["n_containers", "linea_base"] == k.loc["n_containers", "plan_optimizado"] == len(res.containers)


def test_politica_de_parametros():
    firma = list(inspect.signature(build_params).parameters)
    assert len(firma) == 6
    sig_main = set(inspect.signature(main_mod.main).parameters) - {"data_dir", "out_dir", "time_limit_s", "workers", "seed", "verbose"}
    assert len(sig_main) == 6
    src = (ROOT / "main.py").read_text(encoding="utf-8")
    flags = set(re.findall(r'add_argument\("(--[a-z-]+)"', src))
    assert flags == {"--fecha", "--horizonte", "--cap-sofasa", "--cap-suppla-semana", "--cap-traslado", "--dias-traslado"}


def test_equivalencia_lote_manual(corrida):
    """Mismos contenedores y parámetros por la ruta de archivos y por la manual: misma clasificación, base, costos y validaciones."""
    res, _ = corrida
    p = build_params(FECHA, HORIZONTE, CAP_S, CAP_P, CAP_T, DIAS_T)
    df = res.df_universo
    sub = df[df["estado_universo"] != "EXCLUIDO_EN_PLANTA"]  # sin EM real no hay exclusión en manual
    manual = [{"id": r.id, "eta": r.eta, "ruptura": None if pd.isna(r.ruptura) else r.ruptura, "vessel": r.vessel, "family": r.family}
              for r in sub.itertuples()]
    cm, dm = prepare_universe(manual, p)
    clas_a = df.set_index("id").loc[sub["id"], "estado_universo"]
    clas_b = dm.set_index("id").loc[sub["id"], "estado_universo"]
    assert (clas_a == clas_b).all()
    cf = [c for c in res.containers]
    assert sorted((c.id, c.eta, c.need_by, c.em_objetivo) for c in cf) == sorted((c.id, c.eta, c.need_by, c.em_objetivo) for c in cm)
    r2 = run_pipeline(cm, p, dm, time_limit_s=5, workers=1, seed=1)
    assert r2.costos_base.totals == res.costos_base.totals
    assert {k: (v.route, v.em_day, v.suppla_in_day) for k, v in r2.base_plan.items()} == \
           {k: (v.route, v.em_day, v.suppla_in_day) for k, v in res.base_plan.items()}
    assert [v[1] for v in r2.validaciones if "base" in v[0] or "Embudo" not in v[0]].count("FALLA") == 0


def test_weighted_objective_optimal_le_base(corrida):
    res, _ = corrida
    p = res.params
    obj_base = weighted_objective(res.containers, res.base_plan, p)
    obj_opt = weighted_objective(res.containers, res.plan, p)
    assert obj_opt <= obj_base + 1e-5


def test_no_trajectory_in_src():
    src_dir = ROOT / "src"
    for path in src_dir.rglob("*.py"):
        content = path.read_text(encoding="utf-8")
        assert not re.search(r"trajectory|trayectoria", content, re.IGNORECASE), f"Encontrado trajectory/trayectoria en {path}"


def test_app_no_render_validaciones():
    app_text = (ROOT / "app.py").read_text(encoding="utf-8")
    assert "render_validaciones" not in app_text


def test_glossary_content(corrida):
    res, _ = corrida
    g = get_glossary(res.params)
    terms = [item[0] for item in g]
    assert "ETD" in terms
    assert "ETA" in terms
    assert "EM objetivo" in terms
    assert "EM previa" in terms
    all_text = " ".join(str(item) for item in g)
    for val in ["PROGRAMABLE", "EXCLUIDO_EN_PLANTA", "FUERA_VENTANA", "PROGRAMADO_EN_HORIZONTE", "EM_PROYECTADA", "DIFERIDO_SIN_CUPO"]:
        assert val in all_text, f"Falta el valor de estado {val} en el glosario"


def test_print_validaciones_no_falla(corrida, capsys):
    res, _ = corrida
    print_validaciones(res)
    captured = capsys.readouterr().out
    assert "[FALLA]" not in captured
    assert "OK" in captured
